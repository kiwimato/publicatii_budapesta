"""Combină tabelele din research/*.md într-un tabel centralizat (Markdown, CSV, XLSX)."""
import csv, re
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
CATEGORII = {
    "01a_publicatii_online_cultura.md": "1. Publicații online de cultură",
    "01b_presa_scrisa_cultura.md": "1. Presă scrisă de cultură",
    "02_social_media.md": "2. Social media",
    "03_posturi_tv.md": "3. Posturi TV / radio",
    "04_publicatii_romanesti_ungaria.md": "4. Publicații românești din Ungaria",
}
COLS = ["Categorie", "Secțiune", "Nume", "Tip / profil", "Website / link",
        "Persoane de contact", "Email", "Telefon", "Adresă", "Sursă contact", "Observații"]
MAP = {  # coloana comună -> antete posibile din fișierele sursă
    "Nume": ["Publicație", "Nume", "Post TV", "Post radio", "Post TV / emisiune", "Publicație / Instituție media"],
    "Tip / profil": ["Profil", "Tip", "Platformă"],
    "Website / link": ["Website", "Link"],
    "Persoane de contact": ["Persoane de contact (funcție)", "Emisiuni relevante / persoane de contact", "Contact admin / email / telefon"],
    "Email": ["Email"], "Telefon": ["Telefon"], "Adresă": ["Adresă"],
    "Sursă contact": ["Sursă contact", "Sursă/confirmare"], "Observații": ["Observații"],
}
EXTRA = ["Periodicitate", "Limbă", "Editor", "Proprietar", "Membri/urmăritori", "Public/Privat", "Reguli postare"]

def cells(line):
    return [c.strip() for c in re.split(r"(?<!\\)\|", line.strip().strip("|"))]

rows = []
for fname, cat in CATEGORII.items():
    lines = (ROOT / "research" / fname).read_text(encoding="utf-8").splitlines()
    section, header = "", None
    for i, line in enumerate(lines):
        if line.startswith("#"):
            section = line.lstrip("#").strip()
        if not line.startswith("|"):
            header = None
            continue
        if header is None:
            if i + 1 < len(lines) and re.match(r"^\|\s*:?-", lines[i + 1]):
                header = cells(line)
            continue
        if re.match(r"^\|\s*:?-", line) or header[0] != "Nr":
            continue
        rec = dict(zip(header, cells(line)))
        out = {"Categorie": cat, "Secțiune": section}
        for col, keys in MAP.items():
            out[col] = next((rec[k] for k in keys if k in rec), "")
        extra = "; ".join(f"{k}: {rec[k]}" for k in EXTRA if rec.get(k))
        if extra:
            out["Tip / profil"] = f"{out['Tip / profil']} ({extra})" if out["Tip / profil"] else extra
        rows.append(out)

with open(ROOT / "TABEL_CENTRALIZAT.csv", "w", newline="", encoding="utf-8-sig") as f:
    w = csv.DictWriter(f, COLS); w.writeheader(); w.writerows(rows)

md = ["# Tabel centralizat — contacte media ICR Budapesta", "",
      f"Generat din `research/*.md` (date verificate la 2026-10-02). Total: {len(rows)} intrări.", "",
      "| Nr | " + " | ".join(COLS) + " |", "|" + "---|" * (len(COLS) + 1)]
md += [f"| {n} | " + " | ".join(r[c].replace("\n", " ") for c in COLS) + " |" for n, r in enumerate(rows, 1)]
(ROOT / "TABEL_CENTRALIZAT.md").write_text("\n".join(md) + "\n", encoding="utf-8")

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    wb = Workbook(); ws = wb.active; ws.title = "Toate"
    def fill(ws, data):
        ws.append(["Nr"] + COLS)
        for n, r in enumerate(data, 1):
            ws.append([n] + [re.sub(r"\[([^\]]*)\]\(([^)]*)\)", r"\1 (\2)", r[c]).replace("**", "") for c in COLS])
        for c in ws[1]: c.font = Font(bold=True)
        widths = [5, 28, 30, 28, 30, 35, 40, 35, 25, 30, 40, 50]
        for i, wdt in enumerate(widths): ws.column_dimensions[chr(65 + i)].width = wdt
        for row in ws.iter_rows(min_row=2):
            for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "C2"; ws.auto_filter.ref = ws.dimensions
    fill(ws, rows)
    for cat in CATEGORII.values():
        if cat in wb.sheetnames: continue
        sub = [r for r in rows if r["Categorie"] == cat]
        if sub: fill(wb.create_sheet(cat.split(". ", 1)[1].replace("/", "-")[:31]), sub)
    wb.save(ROOT / "TABEL_CENTRALIZAT.xlsx")
except ImportError:
    print("openpyxl lipsă — XLSX nu a fost generat")
print(len(rows), "intrări")
